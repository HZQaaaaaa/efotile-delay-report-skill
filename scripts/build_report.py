# -*- coding: utf-8 -*-
"""
销拓订单延迟率报表构建器（skill 版，无二进制模板依赖）。
输入：平台导出的「销售摘要」xlsx（sheet 名「销售摘要」，51 列）
输出：
  1) 带函数 Excel：销售摘要(新增5列公式) + 店铺名称匹配(内置) + 发货统计—延迟发货
  2) 返回统计结果 dict（供 make_html.py 使用）

用法：
  python build_report.py <下载报表.xlsx> <输出目录> <开始日期YYYY-MM-DD> <结束日期YYYY-MM-DD>
"""
import openpyxl, csv, os, sys, json, datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MATCH_CSV = os.path.join(SKILL_DIR, 'assets', '店铺名称匹配.csv')

CHANNELS = [
    ('太和会', ['普林斯顿', '上海探影', '徐州青葵', '杭州松伽', '广州邦耐', '南京乾璟佳', '宁波尚博']),
    ('新平台', ['安徽中厨', '合肥维森', '北京拾光机', '天津倾心', '宁波紫竹', '徐州青匠']),
    ('其他',   ['同舟益源', '长沙福品', '中科云商', '佛山优优乐']),
]
NEW_HEADERS = ['门店简称', '订单登记日期', '5天未发', '7天未发', '10天未发']
RED = 'C00000'


def load_mapping():
    """读取内置匹配表 CSV（E系统名称,F门店简称,G现有代理商,H分类,I场景）。"""
    rows = []
    with open(MATCH_CSV, encoding='utf-8-sig') as f:
        for row in csv.reader(f):
            if row:
                rows.append(row)
    mapping = {}   # E系统名称 -> G现有代理商（VLOOKUP 第3列）
    for row in rows[1:]:
        if len(row) >= 3 and row[0].strip() and row[2].strip():
            mapping[row[0].strip()] = row[2].strip()
    return rows, mapping


def build(download_path, out_xlsx, start, end):
    start_d = datetime.datetime.strptime(start, '%Y-%m-%d').date()
    end_d = datetime.datetime.strptime(end, '%Y-%m-%d').date()
    today = datetime.date.today()

    wd = openpyxl.load_workbook(download_path)
    dws = wd['销售摘要']
    n_rows = dws.max_row  # 含表头

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '销售摘要'

    # ---- sheet1：销售摘要 ----
    for j, h in enumerate(NEW_HEADERS, start=1):
        ws.cell(row=1, column=j, value=h)
    for c in range(1, 52):  # 下载表 51 列表头 -> F 列起
        ws.cell(row=1, column=c + 5, value=dws.cell(row=1, column=c).value)
    for r in range(2, n_rows + 1):
        for c in range(1, 52):
            src = dws.cell(row=r, column=c)
            dst = ws.cell(row=r, column=c + 5)
            dst.value = src.value
            if src.number_format and src.number_format != 'General':
                dst.number_format = src.number_format
        ws.cell(row=r, column=1).value = '=IFERROR(VLOOKUP(H%d,店铺名称匹配!$E:$G,3,0),"")' % r
        ws.cell(row=r, column=2).value = '=TEXT(AG%d,"yyyy年mm月dd日")' % r
        ws.cell(row=r, column=3).value = '=IF(AND(OR(K%d="待审核",K%d="预订单"),TODAY()-B%d>=5,TODAY()-B%d<7),1,0)' % (r, r, r, r)
        ws.cell(row=r, column=4).value = '=IF(AND(OR(K%d="待审核",K%d="预订单"),TODAY()-B%d>=7,TODAY()-B%d<10),1,0)' % (r, r, r, r)
        ws.cell(row=r, column=5).value = '=IF(AND(OR(K%d="待审核",K%d="预订单"),TODAY()-B%d>=10),1,0)' % (r, r, r)
    last = n_rows

    # ---- sheet：店铺名称匹配（内置，CSV -> E:I 列）----
    mws = wb.create_sheet('店铺名称匹配')
    rows, mapping = load_mapping()
    for ri, row in enumerate(rows, start=1):
        for ci, v in enumerate(row, start=5):  # E=5
            if v != '':
                mws.cell(row=ri, column=ci, value=v)

    # ---- sheet2：发货统计—延迟发货 ----
    s2 = wb.create_sheet('发货统计—延迟发货')
    thin = Side(style='thin', color='D9D9D9')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill('solid', fgColor=RED)
    hdr_font = Font(bold=True, color='FFFFFF')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    s2.merge_cells('A1:A2')
    s2.merge_cells('B1:J1')
    s2['A1'] = '渠道'
    s2['B1'] = '未发货统计'
    if start_d.month == end_d.month:
        lbl_c, lbl_d = '%d月累计发货' % start_d.month, '%d月未发货' % start_d.month
    else:
        lbl_c = '%d-%d月累计发货' % (start_d.month, end_d.month)
        lbl_d = '%d-%d月未发货' % (start_d.month, end_d.month)
    headers2 = ['店铺', lbl_c, lbl_d, '5天未发', '7天未发', '10天未发',
                '5天发货率', '7天发货率', '10天发货率']
    for j, h in enumerate(headers2, start=2):
        s2.cell(row=2, column=j, value=h)

    r = 3
    for ch, shops in CHANNELS:
        start_r = r
        for shop in shops:
            s2.cell(row=r, column=2, value=shop)
            s2.cell(row=r, column=3).value = ('=COUNTIFS(销售摘要!A:A,B%d,销售摘要!K:K,"发货在途单")'
                                              '+COUNTIFS(销售摘要!A:A,B%d,销售摘要!K:K,"待发货")') % (r, r)
            s2.cell(row=r, column=4).value = ('=COUNTIFS(销售摘要!A:A,B%d,销售摘要!K:K,"预订单")'
                                              '+COUNTIFS(销售摘要!A:A,B%d,销售摘要!K:K,"待审核")') % (r, r)
            s2.cell(row=r, column=5).value = '=SUMIF(销售摘要!$A:$A,$B%d,销售摘要!C:C)' % r
            s2.cell(row=r, column=6).value = '=SUMIF(销售摘要!$A:$A,$B%d,销售摘要!D:D)' % r
            s2.cell(row=r, column=7).value = '=SUMIF(销售摘要!$A:$A,$B%d,销售摘要!E:E)' % r
            s2.cell(row=r, column=8).value = f'=IF((C{r}+D{r})=0,"",(C{r}+D{r}-E{r}-F{r}-G{r})/(C{r}+D{r}))'
            s2.cell(row=r, column=9).value = f'=IF((C{r}+D{r})=0,"",(C{r}+D{r}-F{r}-G{r})/(C{r}+D{r}))'
            s2.cell(row=r, column=10).value = f'=IF((C{r}+D{r})=0,"",(C{r}+D{r}-G{r})/(C{r}+D{r}))'
            r += 1
        s2.merge_cells(start_row=start_r, start_column=1, end_row=r - 1, end_column=1)
        cc = s2.cell(row=start_r, column=1, value=ch)
        cc.alignment = center
        cc.font = Font(bold=True)

    s2.merge_cells('A20:B20')
    s2['A20'] = '销售拓展业务合计'
    for col in 'CDEFG':
        s2['%s20' % col].value = '=SUM(%s3:%s19)' % (col, col)
    s2['H20'].value = '=IF((C20+D20)=0,"",(C20+D20-E20-F20-G20)/(C20+D20))'
    s2['I20'].value = '=IF((C20+D20)=0,"",(C20+D20-F20-G20)/(C20+D20))'
    s2['J20'].value = '=IF((C20+D20)=0,"",(C20+D20-G20)/(C20+D20))'

    # 样式
    for coord in ('A1', 'B1') + tuple('%s2' % c for c in 'BCDEFGHIJ'):
        cell = s2[coord]
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = center
    for rr in range(2, 21):
        for cc in range(1, 11):
            cell = s2.cell(row=rr, column=cc)
            cell.border = border
            if cc >= 8:
                cell.number_format = '0%'
            elif cc >= 3 and rr >= 3:
                cell.number_format = '0;\\-0;"-"'
            if cc != 1:
                cell.alignment = Alignment(horizontal='center', vertical='center')
    s2['A20'].font = Font(bold=True)
    s2['A20'].alignment = center
    for col, w in (('A', 12), ('B', 12), ('C', 12), ('D', 12), ('E', 9), ('F', 9),
                   ('G', 9), ('H', 10), ('I', 10), ('J', 10)):
        s2.column_dimensions[col].width = w

    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass
    os.makedirs(os.path.dirname(out_xlsx), exist_ok=True)
    wb.save(out_xlsx)

    # ---- Python 侧同步统计（供 HTML，并作公式独立验算）----
    stats = {}
    for r2 in range(2, n_rows + 1):
        shop_raw = dws.cell(row=r2, column=3).value
        status = dws.cell(row=r2, column=6).value
        reg = dws.cell(row=r2, column=28).value
        shop = mapping.get(str(shop_raw).strip()) if shop_raw else None
        if not shop:
            continue
        st = stats.setdefault(shop, {'C': 0, 'D': 0, 'E': 0, 'F': 0, 'G': 0})
        if status in ('发货在途单', '待发货'):
            st['C'] += 1
        elif status in ('预订单', '待审核'):
            st['D'] += 1
            if isinstance(reg, datetime.datetime):
                age = (today - reg.date()).days
                if 5 <= age < 7:
                    st['E'] += 1
                if 7 <= age < 10:
                    st['F'] += 1
                if age >= 10:
                    st['G'] += 1
    rows_out = []
    for ch, shops in CHANNELS:
        for shop in shops:
            st = stats.get(shop, {'C': 0, 'D': 0, 'E': 0, 'F': 0, 'G': 0})
            tot = st['C'] + st['D']
            rows_out.append({'channel': ch, 'shop': shop, **st,
                             'h': (tot - st['E'] - st['F'] - st['G']) / tot if tot else None,
                             'i': (tot - st['F'] - st['G']) / tot if tot else None,
                             'j': (tot - st['G']) / tot if tot else None})
    total = {k: sum(rw[k] for rw in rows_out) for k in ('C', 'D', 'E', 'F', 'G')}
    ta = total['C'] + total['D']
    total['h'] = (ta - total['E'] - total['F'] - total['G']) / ta if ta else None
    total['i'] = (ta - total['F'] - total['G']) / ta if ta else None
    total['j'] = (ta - total['G']) / ta if ta else None

    return {'start': start, 'end': end, 'label_c': lbl_c, 'label_d': lbl_d,
            'data_rows': n_rows - 1, 'rows': rows_out, 'total': total}


if __name__ == '__main__':
    dl = sys.argv[1]
    out_dir = sys.argv[2]
    start = sys.argv[3]
    end = sys.argv[4]
    base = start.replace('-', '') + '-' + end.replace('-', '')
    out_xlsx = os.path.join(out_dir, '订单延迟率—欣怡_%s.xlsx' % base)
    res = build(dl, out_xlsx, start, end)
    print(json.dumps({'out': out_xlsx, 'data_rows': res['data_rows'],
                      'total': res['total']}, ensure_ascii=False, indent=1))
